import glob
import os
import docx
import fitz

def audit():
    print("=== STARTING EXHAUSTIVE AUDIT ACROSS RESEARCH FOLDER ===")
    
    # 1. Check em-dashes across all text, docx, and pdf files
    em_dash = '\u2014'
    total_em_dashes = 0
    checked_files = 0
    
    all_files = sorted(glob.glob('research/**/*', recursive=True))
    for f in all_files:
        if not os.path.isfile(f):
            continue
        rel = os.path.relpath(f, 'research')
        
        # Text files
        if f.endswith(('.tex', '.md', '.py', '.txt', '.json', '.bib')):
            checked_files += 1
            with open(f, 'r', encoding='utf-8', errors='ignore') as fh:
                c = fh.read()
            cnt = c.count(em_dash)
            if cnt > 0:
                print(f"[FAIL] {rel}: {cnt} em-dashes found!")
                total_em_dashes += cnt
            else:
                pass
                
        # Docx files
        elif f.endswith('.docx') and not os.path.basename(f).startswith('~$'):
            checked_files += 1
            d = docx.Document(f)
            cnt = 0
            for p in d.paragraphs:
                cnt += p.text.count(em_dash)
            for t in d.tables:
                for row in t.rows:
                    for cell in row.cells:
                        cnt += cell.text.count(em_dash)
            if cnt > 0:
                print(f"[FAIL] {rel}: {cnt} em-dashes found!")
                total_em_dashes += cnt
            else:
                pass
                
        # PDF files
        elif f.endswith('.pdf'):
            checked_files += 1
            doc_pdf = fitz.open(f)
            cnt = 0
            for page in doc_pdf:
                cnt += page.get_text().count(em_dash)
            if cnt > 0:
                print(f"[FAIL] {rel}: {cnt} em-dashes found!")
                total_em_dashes += cnt
            else:
                pass

    print(f"Total files audited: {checked_files}")
    print(f"Total em-dashes detected: {total_em_dashes}")
    assert total_em_dashes == 0, "Em-dashes still exist!"

    # 2. Check RQ presence across core docs
    core_docs = {
        "IEEE Paper (LaTeX)": "research/ieee_paper.tex",
        "IEEE Paper (Markdown)": "research/IEEE_Paper_Submission_Manuscript.md",
        "IEEE Paper (DOCX)": "research/IEEE_Research_Paper_Digital_Mine.docx",
        "Conference Presentation": "research/conference_presentation.md",
        "Full Monograph (DOCX)": "research/full_research_paper.docx",
        "Technical Report (DOCX)": "research/technical_report.docx",
        "IEEE PDF": "research/IEEE_Research_Paper_Digital_Mine.pdf",
        "Monograph PDF": "research/Full_Research_Monograph_Digital_Mine.pdf",
        "Presentation PDF": "research/Digital_Mine_Conference_Presentation.pdf"
    }

    print("\n=== CHECKING RESEARCH QUESTIONS (RQ1 - RQ4) PRESENCE ===")
    for name, path in core_docs.items():
        has_rq1 = False
        has_rq4 = False
        if path.endswith(('.tex', '.md')):
            with open(path, 'r', encoding='utf-8') as fh:
                t = fh.read()
            has_rq1 = 'RQ1' in t
            has_rq4 = 'RQ4' in t
        elif path.endswith('.docx'):
            d = docx.Document(path)
            t = '\n'.join([p.text for p in d.paragraphs] + [c.text for tab in d.tables for r in tab.rows for c in r.cells])
            has_rq1 = 'RQ1' in t
            has_rq4 = 'RQ4' in t
        elif path.endswith('.pdf'):
            doc_pdf = fitz.open(path)
            t = '\n'.join([p.get_text() for p in doc_pdf])
            has_rq1 = 'RQ1' in t
            has_rq4 = 'RQ4' in t
        print(f"{name} ({os.path.basename(path)}): RQ1={has_rq1}, RQ4={has_rq4}")
        assert has_rq1 and has_rq4, f"{name} is missing Research Questions!"

    # 3. Check image embedding in PDFs and DOCX
    print("\n=== CHECKING EMBEDDED FIGURES & DIAGRAMS ===")
    d_docx = docx.Document('research/full_research_paper.docx')
    docx_imgs = [r for r in d_docx.part.rels.values() if 'image' in r.target_ref]
    print(f"full_research_paper.docx: {len(docx_imgs)} embedded images")
    assert len(docx_imgs) >= 16, "Not all 16 figures embedded in docx!"

    ieee_docx = docx.Document('research/IEEE_Research_Paper_Digital_Mine.docx')
    ieee_docx_imgs = [r for r in ieee_docx.part.rels.values() if 'image' in r.target_ref]
    print(f"IEEE_Research_Paper_Digital_Mine.docx: {len(ieee_docx_imgs)} embedded images")
    assert len(ieee_docx_imgs) >= 11, f"Expected at least 11 figures in IEEE docx, got {len(ieee_docx_imgs)}!"

    ieee_pdf = fitz.open('research/IEEE_Research_Paper_Digital_Mine.pdf')
    ieee_imgs = sum(len(p.get_images()) for p in ieee_pdf)
    print(f"IEEE Research Paper PDF: {len(ieee_pdf)} pages, {ieee_imgs} embedded images")
    assert ieee_imgs >= 11, f"Expected at least 11 figures in IEEE PDF, got {ieee_imgs}!"

    mono_pdf = fitz.open('research/Full_Research_Monograph_Digital_Mine.pdf')
    mono_imgs = sum(len(p.get_images()) for p in mono_pdf)
    print(f"Full Research Monograph PDF: {len(mono_pdf)} pages, {mono_imgs} embedded images")
    assert mono_imgs >= 16, "Not all 16 figures embedded in Monograph PDF!"

    pres_pdf = fitz.open('research/Digital_Mine_Conference_Presentation.pdf')
    pres_imgs = sum(len(p.get_images()) for p in pres_pdf)
    print(f"Conference Presentation PDF: {len(pres_pdf)} slides, {pres_imgs} embedded images")
    assert pres_imgs >= 6, "Not all 6 figures embedded in Presentation PDF!"

    print("\n=== ALL AUDIT CHECKS PASSED PERFECTLY ===")

if __name__ == '__main__':
    audit()
