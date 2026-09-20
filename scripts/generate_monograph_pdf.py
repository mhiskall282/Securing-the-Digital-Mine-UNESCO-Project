"""Generate a publication-grade, fully illustrated PDF for the complete research monograph directly from research/full_research_paper.docx with embedded figures, tables, and running page numbers."""

import os
import docx
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, Image, HRFlowable, KeepTogether
)
from reportlab.pdfgen import canvas

class MonographNumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super().showPage()
        super().save()

    def draw_page_decorations(self, total_pages):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#555555"))
        
        # Header (pages > 1)
        if self._pageNumber > 1:
            self.drawString(54, 750, "UNESCO RUSSIAN-AFRICAN FORUM 2026  |  DESIGN SCIENCE RESEARCH MONOGRAPH")
            self.setStrokeColor(colors.HexColor("#D0D7DE"))
            self.setLineWidth(0.5)
            self.line(54, 744, 558, 744)

        # Footer (all pages)
        footer_text = f"Securing the Digital Mine: DSR Research Monograph   |   Page {self._pageNumber} of {total_pages}"
        self.drawRightString(558, 36, footer_text)
        self.setStrokeColor(colors.HexColor("#D0D7DE"))
        self.setLineWidth(0.5)
        self.line(54, 46, 558, 46)
        
        self.restoreState()


def build_monograph_from_docx(docx_path: str, output_path: str):
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc_tpl = SimpleDocTemplate(
        output_path,
        pagesize=letter,
        leftMargin=54,
        rightMargin=54,
        topMargin=60,
        bottomMargin=56
    )

    styles = getSampleStyleSheet()
    c_navy = colors.HexColor("#0B2545")
    c_blue = colors.HexColor("#134074")
    c_slate = colors.HexColor("#1D2D44")
    c_accent = colors.HexColor("#8DA9C4")
    c_body = colors.HexColor("#222222")
    c_tint = colors.HexColor("#F0F4F8")
    c_code_bg = colors.HexColor("#F8FAFC")

    styles.add(ParagraphStyle('MonoTitle', fontName='Helvetica-Bold', fontSize=18, leading=22, textColor=c_navy, alignment=1, spaceAfter=8))
    styles.add(ParagraphStyle('MonoSubtitle', fontName='Helvetica', fontSize=10, leading=14, textColor=c_blue, alignment=1, spaceAfter=12))
    styles.add(ParagraphStyle('ChapterHead', fontName='Helvetica-Bold', fontSize=13, leading=17, textColor=c_navy, spaceBefore=14, spaceAfter=6, keepWithNext=True))
    styles.add(ParagraphStyle('SubHead', fontName='Helvetica-Bold', fontSize=10.5, leading=14, textColor=c_blue, spaceBefore=10, spaceAfter=4, keepWithNext=True))
    styles.add(ParagraphStyle('SubSubHead', fontName='Helvetica-Bold', fontSize=9.5, leading=13, textColor=c_slate, spaceBefore=8, spaceAfter=3, keepWithNext=True))
    styles.add(ParagraphStyle('MonoBody', fontName='Helvetica', fontSize=9, leading=13, textColor=c_body, spaceAfter=5, alignment=4))
    styles.add(ParagraphStyle('MonoBullet', fontName='Helvetica', fontSize=9, leading=13, textColor=c_body, leftIndent=14, firstLineIndent=-10, spaceAfter=3))
    styles.add(ParagraphStyle('MonoCode', fontName='Courier', fontSize=8, leading=11, textColor=c_slate, spaceBefore=3, spaceAfter=4))
    styles.add(ParagraphStyle('FigureCaption', fontName='Helvetica-Oblique', fontSize=8.5, leading=11, textColor=colors.HexColor("#444444"), alignment=1, spaceBefore=4, spaceAfter=8))
    styles.add(ParagraphStyle('TableHead', fontName='Helvetica-Bold', fontSize=8, leading=10.5, textColor=colors.white, alignment=1))
    styles.add(ParagraphStyle('TableCell', fontName='Helvetica', fontSize=7.5, leading=9.5, textColor=c_body))

    # Figure mapping
    fig_map = {
        "Figure 3.1: Six-Stage Design Science Research Process Framework": ("research/figures/dsr_framework.png", 460, 190),
        "Figure 3.2: Four-Layer End-to-End System Architecture": ("research/figures/system_architecture.png", 460, 270),
        "Figure 3.3: Cyber-Physical Mineral Processing SCADA Circuit": ("research/figures/mining_scada_flowchart.png", 460, 228),
        "Figure 3.4: Database Entity-Relationship (ER) Schema": ("research/figures/er_diagram.png", 440, 280),
        "Figure 3.5: UML Use Case Diagram": ("research/figures/uml_use_case.png", 420, 260),
        "Figure 3.6: UML Class Diagram": ("research/figures/uml_class_diagram.png", 420, 270),
        "Figure 3.7: UML Activity Diagram": ("research/figures/uml_activity_diagram.png", 420, 280),
        "Figure 3.8: UML Sequence Diagram": ("research/figures/uml_sequence_diagram.png", 420, 260),
        "Figure 3.9: Interface Design Wireframe": ("research/figures/dashboard_wireframe.png", 460, 270),
        "Figure 3.10: Spatial-Temporal CNN-LSTM Deep Neural Network Flowchart": ("research/figures/cnn_lstm_architecture.png", 460, 228),
        "Figure 4.1: BWOA Fitness Convergence History across 100 Iterations": ("research/figures/bwoa_convergence.png", 440, 275),
        "Figure 4.2: Gini Feature Importance Ranking (Selected 10 vs": ("research/figures/feature_importance.png", 440, 264),
        "Figure 4.3: CNN-LSTM Loss and Accuracy Convergence History": ("research/figures/training_curves.png", 420, 260),
        "Figure 4.4: Confusion Matrix on Held-Out KDDTest+ Benchmark": ("research/figures/confusion_matrix.png", 280, 280),
        "Figure 4.5: Receiver Operating Characteristic (ROC) Curves": ("research/figures/roc_auc_curves.png", 340, 255),
        "Figure 4.6: Single-Sample Inference Latency vs SCADA Real-Time Ceiling": ("research/figures/latency_comparison_barchart.png", 440, 245),
    }

    doc = docx.Document(docx_path)
    story = []

    # Title & Metadata
    if len(doc.paragraphs) > 0:
        story.append(Paragraph("Securing the Digital Mine: Edge Intrusion Detection in Industrial Mining IoT", styles['MonoTitle']))
        story.append(Paragraph("A Design Science Research Monograph on Metaheuristic-Optimized Deep Learning for Operational SCADA Security", styles['MonoSubtitle']))
        story.append(HRFlowable(width="100%", thickness=1, color=c_navy, spaceBefore=0, spaceAfter=10))

    # Process paragraphs
    for p_idx, p in enumerate(doc.paragraphs):
        p_text = p.text.strip()
        if not p_text:
            continue
            
        # Check for Chapter heading
        if p_text.startswith("CHAPTER") or p_text.startswith("APPENDICES"):
            story.append(PageBreak())
            story.append(Paragraph(f"<b>{p_text}</b>", styles['ChapterHead']))
            story.append(HRFlowable(width="100%", thickness=0.75, color=c_blue, spaceBefore=2, spaceAfter=8))
            continue
            
        # Check for Section / Subsection heading
        if any(p_text.startswith(f"{i}.") for i in range(1, 10)) and len(p_text) < 80:
            story.append(Paragraph(f"<b>{p_text}</b>", styles['SubHead']))
            continue
            
        if p_text.startswith("APPENDIX") and len(p_text) < 90:
            story.append(Spacer(1, 6))
            story.append(Paragraph(f"<b>{p_text}</b>", styles['SubHead']))
            story.append(HRFlowable(width="100%", thickness=0.5, color=c_accent, spaceBefore=1, spaceAfter=6))
            continue

        # Check for Figure Caption
        is_fig = False
        for fig_title, (img_path, w, h) in fig_map.items():
            if fig_title in p_text and os.path.exists(img_path):
                story.append(Spacer(1, 4))
                story.append(Image(img_path, width=w, height=h))
                story.append(Spacer(1, 2))
                story.append(Paragraph(p_text, styles['FigureCaption']))
                story.append(Spacer(1, 6))
                is_fig = True
                break
        if is_fig:
            continue

        # Check for bullet points
        if p_text.startswith("• ") or p_text.startswith("* ") or p_text.startswith("- ") or (len(p_text) > 3 and p_text[0].isdigit() and p_text[1] in ['.', ')']):
            story.append(Paragraph(p_text, styles['MonoBullet']))
            continue

        # Check for Equation / Algorithm blocks
        if "Equation " in p_text or "Algorithm 1:" in p_text or "Layer 1:" in p_text:
            # Format as monospace callout
            code_lines = p_text.split('\n')
            formatted_code = "<br/>".join(line.replace(" ", "&nbsp;") for line in code_lines)
            code_table = Table([[Paragraph(formatted_code, styles['MonoCode'])]], colWidths=[504])
            code_table.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,-1), c_code_bg),
                ('BOX', (0,0), (-1,-1), 0.5, c_accent),
                ('PADDING', (0,0), (-1,-1), 6),
            ]))
            story.append(Spacer(1, 3))
            story.append(code_table)
            story.append(Spacer(1, 4))
            continue

        # Regular academic paragraph
        story.append(Paragraph(p_text, styles['MonoBody']))

    # Also render tables from docx
    if doc.tables:
        story.append(PageBreak())
        story.append(Paragraph("<b>COMPREHENSIVE RESEARCH MONOGRAPH DATA TABLES</b>", styles['ChapterHead']))
        story.append(HRFlowable(width="100%", thickness=0.75, color=c_blue, spaceBefore=2, spaceAfter=8))

        for t_idx, t in enumerate(doc.tables):
            if len(t.rows) < 2:
                continue
            table_data = []
            for r_idx, row in enumerate(t.rows):
                row_data = []
                for cell in row.cells:
                    cell_text = cell.text.strip().replace('\n', ' ')
                    if r_idx == 0:
                        row_data.append(Paragraph(f"<b>{cell_text}</b>", styles['TableHead']))
                    else:
                        row_data.append(Paragraph(cell_text, styles['TableCell']))
                table_data.append(row_data)

            num_cols = len(t.columns)
            col_w = 504 / num_cols
            rl_table = Table(table_data, colWidths=[col_w]*num_cols)
            rl_table.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,0), c_navy),
                ('ALIGN', (0,0), (-1,-1), 'CENTER'),
                ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
                ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#CCCCCC")),
                ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, c_tint]),
                ('PADDING', (0,0), (-1,-1), 4),
            ]))
            story.append(Spacer(1, 6))
            story.append(rl_table)
            story.append(Spacer(1, 8))

    doc_tpl.build(story, canvasmaker=MonographNumberedCanvas)
    print(f"Monograph PDF successfully compiled at: {output_path}")

if __name__ == '__main__':
    docx_file = os.path.join("research", "full_research_paper.docx")
    out_pdf = os.path.join("research", "Full_Research_Monograph_Digital_Mine.pdf")
    build_monograph_from_docx(docx_file, out_pdf)
