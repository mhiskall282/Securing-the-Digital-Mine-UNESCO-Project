import docx

doc = docx.Document('research/full_research_paper.docx')
print(f"Total paragraphs: {len(doc.paragraphs)}")
for i, p in enumerate(doc.paragraphs):
    if not p.text.strip():
        prev_text = doc.paragraphs[i-1].text.strip() if i > 0 else 'START'
        next_text = doc.paragraphs[i+1].text.strip() if i < len(doc.paragraphs)-1 else 'END'
        print(f"Empty P{i}: after '{prev_text[:60]}' before '{next_text[:60]}'")
